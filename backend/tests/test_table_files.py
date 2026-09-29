import csv
import io
import json

import pytest
from openpyxl import load_workbook

from app.domain.table_files import write_table


def test_csv_starts_with_bom_and_keeps_unicode() -> None:
    content = write_table(
        ["City", "Amount"],
        [{"City": "Москва", "Amount": None}],
        "csv",
    )

    assert content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows == [["City", "Amount"], ["Москва", ""]]


def test_json_array_length_and_null() -> None:
    content = write_table(
        ["City", "Amount"],
        [
            {"Amount": None, "City": "Москва"},
            {"City": "London", "Amount": 2},
        ],
        "json",
    )

    text = content.decode("utf-8")
    assert "Москва" in text
    rows = json.loads(text)
    assert len(rows) == 2
    assert rows[0]["Amount"] is None
    assert list(rows[0]) == ["City", "Amount"]
    assert rows[1] == {"City": "London", "Amount": 2}


def test_xlsx_sheet_is_named_data() -> None:
    content = write_table(
        ["City", "Amount"],
        [{"City": "Москва", "Amount": None}],
        "xlsx",
    )

    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Data"]
    sheet = workbook["Data"]
    assert [cell.value for cell in sheet[1]] == ["City", "Amount"]
    assert sheet["A2"].value == "Москва"
    assert sheet["B2"].value is None


def test_unknown_format_raises() -> None:
    with pytest.raises(ValueError):
        write_table(["City"], [{"City": "Москва"}], "pdf")
