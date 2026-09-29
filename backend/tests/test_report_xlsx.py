from io import BytesIO

from openpyxl import load_workbook

from app.domain.report_model import ReportBlock, ReportTable
from app.domain.report_xlsx import report_xlsx

_SHEETS = (
    "Preview",
    "Columns",
    "Shape",
    "Types",
    "Missing",
    "Summary",
    "Ranking",
    "Grouping",
    "Timeseries",
    "Insights",
)

_INSIGHT = "PnL decreased by 30.1%"


def _table(headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...]) -> ReportTable:
    return ReportTable(headers=headers, rows=rows)


def _blocks() -> tuple[ReportBlock, ...]:
    return (
        ReportBlock(
            "Preview",
            "Preview",
            None,
            (_table(("Name", "Value"), (("alpha", "1"),)),),
        ),
        ReportBlock(
            "Columns",
            "Columns",
            None,
            (_table(("name", "dtype"), (("Name", "text"),)),),
        ),
        ReportBlock(
            "Shape",
            "Size",
            None,
            (_table(("row_count", "column_count"), (("2", "2"),)),),
        ),
        ReportBlock(
            "Types",
            "Data types",
            None,
            (_table(("name", "dtype"), (("Name", "text"),)),),
        ),
        ReportBlock(
            "Missing",
            "Missing values",
            None,
            (_table(("name", "missing_count"), (("Name", "0"),)),),
        ),
        ReportBlock(
            "Summary",
            "Summary",
            None,
            (
                _table(("column", "mean"), (("Value", "1"),)),
                _table(("column", "role"), (("Name", "category"),)),
            ),
        ),
        ReportBlock(
            "Ranking",
            "Top and worst",
            None,
            (
                _table(("Top", "rank"), (("Top", "1"),)),
                _table(("Worst", "rank"), (("Worst", "1"),)),
            ),
        ),
        ReportBlock(
            "Grouping",
            "Grouping",
            None,
            (_table(("value", "count"), (("east", "1"),)),),
        ),
        ReportBlock("Timeseries", "Dynamics", "No date column.", ()),
        ReportBlock(
            "Insights",
            "Insights",
            None,
            (_table(("text",), ((_INSIGHT,),)),),
        ),
    )


def test_sheets_follow_block_order_and_insight_text_is_readable() -> None:
    content = report_xlsx(_blocks())
    workbook = load_workbook(BytesIO(content))

    assert workbook.sheetnames == list(_SHEETS)
    insights = workbook["Insights"]
    assert insights["A1"].value == "Insights"
    assert insights["A2"].value == "text"
    assert insights["A3"].value == _INSIGHT

    summary = workbook["Summary"]
    assert summary["A1"].value == "Summary"
    assert summary["A2"].value == "column"
    assert summary["A3"].value == "Value"
    assert summary["A4"].value is None
    assert summary["A5"].value == "column"
    assert summary["B5"].value == "role"


def test_message_only_block_does_not_raise() -> None:
    block = ReportBlock("Timeseries", "Dynamics", "No date column.", ())

    content = report_xlsx((block,))

    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Timeseries"]
    assert sheet["A1"].value == "Dynamics"
    assert sheet["A2"].value == "No date column."
    assert sheet["A3"].value is None
