import csv
import io
import json
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.domain.messages import message
from app.infrastructure.store import DatasetStore, set_store
from app.main import app

client = TestClient(app)

_SHEETS = [
    "Preview",
    "Columns",
    "Shape",
    "Types",
    "Missing",
    "Summary",
    "Ranking",
    "Grouping",
    "Timeseries",
    "Insights",
]


def _sample_workbook() -> Path:
    container = Path("/data/global_brokerage_dataset.xlsx")
    if container.exists():
        return container
    return Path(__file__).resolve().parents[2] / "data" / "global_brokerage_dataset.xlsx"


class _Clock:
    def __init__(self, now: datetime) -> None:
        self.now = now

    def __call__(self) -> datetime:
        return self.now


def setup_function() -> None:
    holder = _Clock(datetime.now(timezone.utc).replace(microsecond=0))
    set_store(DatasetStore(ttl=timedelta(minutes=60), clock=holder))


def teardown_function() -> None:
    set_store(None)


def _upload(name: str, content: bytes) -> str:
    response = client.post("/api/v1/datasets", files={"file": (name, content)})
    assert response.status_code == 201
    created = response.json()
    assert created["status"] == "ready"
    return created["dataset_id"]


def _upload_sample() -> str:
    workbook = _sample_workbook()
    return _upload(workbook.name, workbook.read_bytes())


def _wide_csv() -> bytes:
    lines = ["Region,City,Amount"]
    for index in range(150):
        region = "keep" if index < 120 else "drop"
        city = "Москва" if index == 0 else "London"
        lines.append(f"{region},{city},{index}")
    return ("\n".join(lines) + "\n").encode()


def _keep_body() -> dict:
    return {
        "filter": {
            "combinator": "and",
            "conditions": [{"column": "Region", "operator": "eq", "value": "keep"}],
        },
        "page": 9,
    }


def _sheet_text(content: bytes) -> dict[str, str]:
    workbook = load_workbook(io.BytesIO(content))
    found: dict[str, str] = {}
    for name in workbook.sheetnames:
        cells = []
        for row in workbook[name].iter_rows(values_only=True):
            cells.extend("" if value is None else str(value) for value in row)
        found[name] = "\n".join(cells)
    return found


def _pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def test_sample_report_xlsx_has_ten_sheets_and_both_languages() -> None:
    dataset_id = _upload_sample()

    english = client.get(
        f"/api/v1/datasets/{dataset_id}/report",
        params={"format": "xlsx", "lang": "en"},
    )
    assert english.status_code == 200
    assert english.headers["content-disposition"] == (
        'attachment; filename="global_brokerage_dataset.xlsx"'
    )
    english_text = _sheet_text(english.content)
    assert list(english_text) == _SHEETS
    assert "PnL decreased by" in english_text["Insights"]
    assert "2025-01-01" in english_text["Insights"]

    russian = client.get(
        f"/api/v1/datasets/{dataset_id}/report",
        params={"format": "xlsx", "lang": "ru"},
    )
    assert russian.status_code == 200
    russian_text = _sheet_text(russian.content)
    assert list(russian_text) == _SHEETS
    assert "PnL снизился на" in russian_text["Insights"]
    assert "Инсайты" in russian_text["Insights"]


def test_sample_report_pdf_contains_the_insight_sentence() -> None:
    dataset_id = _upload_sample()

    english = client.get(
        f"/api/v1/datasets/{dataset_id}/report",
        params={"format": "pdf", "lang": "en"},
    )
    russian = client.get(
        f"/api/v1/datasets/{dataset_id}/report",
        params={"format": "pdf", "lang": "ru"},
    )

    assert english.status_code == 200
    assert 'filename="global_brokerage_dataset.pdf"' in english.headers["content-disposition"]
    assert "PnL decreased by" in _pdf_text(english.content)
    assert russian.status_code == 200
    assert "PnL снизился на" in _pdf_text(russian.content)


def test_unavailable_section_still_downloads() -> None:
    dataset_id = _upload("plain.csv", b"name,amount\nalice,10\nbob,20\n")

    response = client.get(
        f"/api/v1/datasets/{dataset_id}/report",
        params={"format": "xlsx", "lang": "en"},
    )

    assert response.status_code == 200
    sheets = _sheet_text(response.content)
    assert list(sheets) == _SHEETS
    assert sheets["Timeseries"].split("\n")[1] == message("no_datetime_column", "en")


def test_filtered_csv_matches_total_rows_past_the_first_page() -> None:
    dataset_id = _upload("wide.csv", _wide_csv())
    body = _keep_body()

    queried = client.post(f"/api/v1/datasets/{dataset_id}/rows", json=body)
    exported = client.post(
        f"/api/v1/datasets/{dataset_id}/rows/export",
        params={"format": "csv"},
        json=body,
    )

    assert queried.status_code == 422
    assert queried.json()["error"]["code"] == "invalid_page"
    assert exported.status_code == 200
    assert exported.content.startswith(b"\xef\xbb\xbf")
    assert 'filename="wide.csv"' in exported.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
    assert rows[0] == ["Region", "City", "Amount"]
    assert rows[1][1] == "Москва"
    assert len(rows) - 1 == 120

    counted = client.post(
        f"/api/v1/datasets/{dataset_id}/rows",
        json={**body, "page": 1},
    )
    assert counted.status_code == 200
    assert counted.json()["total_rows"] == 120
    assert len(counted.json()["rows"]) == 100


def test_json_and_xlsx_exports_match_and_aggregate_is_groups() -> None:
    dataset_id = _upload("wide.csv", _wide_csv())
    body = _keep_body()

    as_json = client.post(
        f"/api/v1/datasets/{dataset_id}/rows/export",
        params={"format": "json", "lang": "ru"},
        json=body,
    )
    as_xlsx = client.post(
        f"/api/v1/datasets/{dataset_id}/rows/export",
        params={"format": "xlsx"},
        json=body,
    )

    assert as_json.status_code == 200
    payload = json.loads(as_json.content)
    assert isinstance(payload, list)
    assert len(payload) == 120
    assert payload[0]["City"] == "Москва"

    workbook = load_workbook(io.BytesIO(as_xlsx.content))
    assert workbook.sheetnames == ["Data"]
    assert workbook["Data"].max_row == 121

    grouped = client.post(
        f"/api/v1/datasets/{dataset_id}/rows/export",
        params={"format": "json"},
        json={"group": {"mode": "aggregate", "columns": ["Region"]}},
    )
    assert grouped.status_code == 200
    groups = json.loads(grouped.content)
    assert len(groups) == 2
    assert {row["Region"] for row in groups} == {"keep", "drop"}
    assert "row_count" in groups[0]
    assert "Amount" not in groups[0]


def test_missing_and_unknown_export_formats_are_refused() -> None:
    dataset_id = _upload_sample()

    missing = client.get(f"/api/v1/datasets/{dataset_id}/report")
    unknown = client.get(
        f"/api/v1/datasets/{dataset_id}/report",
        params={"format": "csv"},
    )
    rows_unknown = client.post(
        f"/api/v1/datasets/{dataset_id}/rows/export",
        params={"format": "pdf"},
        json={},
    )
    bad_lang = client.get(
        f"/api/v1/datasets/{dataset_id}/report",
        params={"format": "xlsx", "lang": "de"},
    )
    missing_dataset = client.get(
        f"/api/v1/datasets/{uuid.uuid4()}/report",
        params={"format": "xlsx"},
    )

    assert missing.status_code == 422
    assert missing.json()["error"]["code"] == "unsupported_export_format"
    assert missing.json()["error"]["message"] == message(
        "unsupported_export_format", "en"
    )
    assert unknown.status_code == 422
    assert unknown.json()["error"]["code"] == "unsupported_export_format"
    assert rows_unknown.status_code == 422
    assert rows_unknown.json()["error"]["code"] == "unsupported_export_format"
    assert bad_lang.status_code == 422
    assert bad_lang.json()["error"]["code"] == "invalid_language"
    assert missing_dataset.status_code == 404
    assert missing_dataset.json()["error"]["code"] == "dataset_not_found"
