"""Read one uploaded table into a pandas DataFrame.

Supported extensions are csv, tsv, xlsx, xls, json, and parquet.
The first row of a text or Excel file is the header.
"""

import csv
import io
import json
from pathlib import Path

import pandas as pd
import xlrd
from charset_normalizer import from_bytes
from openpyxl import load_workbook

_ALLOWED = {".csv", ".tsv", ".xlsx", ".xls", ".json", ".parquet"}
_DELIMITERS = ",\t;|"


class ReadRejected(Exception):
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def sheet_names(content: bytes, filename: str) -> list[str]:
    """Worksheet names for .xlsx (openpyxl) and .xls (xlrd).

    Raise ReadRejected('unreadable_file') if the workbook cannot be opened.
    """

    extension = _extension(filename)
    try:
        if extension == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                return list(workbook.sheetnames)
            finally:
                workbook.close()
        if extension == ".xls":
            book = xlrd.open_workbook(file_contents=content)
            return list(book.sheet_names())
        raise ReadRejected("unreadable_file")
    except ReadRejected:
        raise
    except Exception as exc:
        raise ReadRejected("unreadable_file") from exc


def read_table(content: bytes, filename: str, sheet: str | None) -> pd.DataFrame:
    """Parse one table. `sheet` is used only for Excel.

    Raise ReadRejected with code: unsupported_format, unreadable_file,
    encoding_error, not_a_table, no_data_rows, sheet_not_found.
    """

    extension = _extension(filename)
    if extension not in _ALLOWED:
        raise ReadRejected("unsupported_format")
    try:
        if extension in {".csv", ".tsv"}:
            return _read_delimited(content, extension)
        if extension == ".json":
            return _read_json(content)
        if extension == ".parquet":
            return _read_parquet(content)
        return _read_excel(content, filename, extension, sheet)
    except ReadRejected:
        raise
    except Exception as exc:
        raise ReadRejected("unreadable_file") from exc


def _extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def _detect_encoding(content: bytes) -> str:
    match = from_bytes(content).best()
    if match is None:
        raise ReadRejected("encoding_error")
    return match.encoding


def _open_text(content: bytes) -> tuple[str, str]:
    encoding = _detect_encoding(content)
    if content.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    try:
        text = content.decode(encoding)
    except (LookupError, UnicodeDecodeError) as exc:
        raise ReadRejected("unreadable_file") from exc
    if text.startswith("\ufeff"):
        text = text.lstrip("\ufeff")
    return text, encoding


def _delimiter(text: str, extension: str) -> str:
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=_DELIMITERS)
    except csv.Error:
        return "," if extension == ".csv" else "\t"
    return dialect.delimiter


def _read_delimited(content: bytes, extension: str) -> pd.DataFrame:
    text, encoding = _open_text(content)
    delimiter = _delimiter(text, extension)
    header = _delimited_header(text, delimiter)
    frame = pd.read_csv(
        io.BytesIO(content),
        encoding=encoding,
        sep=delimiter,
        header=0,
        skip_blank_lines=False,
    )
    return _frame_with_header(frame, header)


def _delimited_header(text: str, delimiter: str) -> list[object]:
    if text == "":
        raise ReadRejected("not_a_table")
    try:
        row = next(csv.reader(io.StringIO(text), delimiter=delimiter))
    except StopIteration:
        raise ReadRejected("not_a_table")
    except csv.Error as exc:
        raise ReadRejected("unreadable_file") from exc
    return list(row)


def _read_json(content: bytes) -> pd.DataFrame:
    text, _encoding = _open_text(content)
    stripped = text.strip()
    if stripped == "":
        raise ReadRejected("not_a_table")
    if stripped.startswith("["):
        try:
            payload = json.loads(stripped)
        except json.JSONDecodeError as exc:
            raise ReadRejected("unreadable_file") from exc
        if not isinstance(payload, list) or any(not isinstance(item, dict) for item in payload):
            raise ReadRejected("not_a_table")
        return _frame_from_named(pd.DataFrame(payload))
    try:
        whole = json.loads(stripped)
    except json.JSONDecodeError:
        whole = None
    if isinstance(whole, dict):
        raise ReadRejected("not_a_table")
    if isinstance(whole, list):
        raise ReadRejected("not_a_table")
    return _frame_from_named(pd.DataFrame(_ndjson_records(text)))


def _ndjson_records(text: str) -> list[dict]:
    records: list[dict] = []
    for line in text.splitlines():
        if line.strip() == "":
            continue
        try:
            item = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ReadRejected("unreadable_file") from exc
        if not isinstance(item, dict):
            raise ReadRejected("not_a_table")
        records.append(item)
    if not records:
        raise ReadRejected("not_a_table")
    return records


def _read_parquet(content: bytes) -> pd.DataFrame:
    frame = pd.read_parquet(io.BytesIO(content), engine="pyarrow")
    return _frame_from_named(frame)


def _read_excel(
    content: bytes,
    filename: str,
    extension: str,
    sheet: str | None,
) -> pd.DataFrame:
    names = sheet_names(content, filename)
    if sheet not in names:
        raise ReadRejected("sheet_not_found")
    engine = "openpyxl" if extension == ".xlsx" else "xlrd"
    frame = pd.read_excel(
        io.BytesIO(content),
        sheet_name=sheet,
        engine=engine,
        header=0,
    )
    header = _excel_header(content, extension, sheet)
    return _frame_with_header(frame, header)


def _excel_header(content: bytes, extension: str, sheet: str) -> list[object]:
    if extension == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        try:
            worksheet = workbook[sheet]
            row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if row is None:
                return []
            return list(row)
        finally:
            workbook.close()
    book = xlrd.open_workbook(file_contents=content)
    worksheet = book.sheet_by_name(sheet)
    if worksheet.nrows == 0 or worksheet.ncols == 0:
        return []
    return [worksheet.cell_value(0, column) for column in range(worksheet.ncols)]


def _frame_with_header(frame: pd.DataFrame, header: list[object]) -> pd.DataFrame:
    if frame.shape[1] == 0:
        raise ReadRejected("not_a_table")
    header = _fit_header(header, frame.shape[1])
    columns = _validated_columns(header)
    return _apply_columns(frame, columns)


def _frame_from_named(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.shape[1] == 0:
        raise ReadRejected("not_a_table")
    columns = _validated_columns(list(frame.columns))
    return _apply_columns(frame, columns)


def _fit_header(header: list[object], width: int) -> list[object]:
    cells = list(header)
    while len(cells) > width and _is_blank_header(cells[-1]):
        cells.pop()
    if len(cells) < width:
        cells.extend([None] * (width - len(cells)))
    if len(cells) != width:
        raise ReadRejected("unreadable_file")
    return cells


def _validated_columns(header: list[object]) -> list[str]:
    if len(header) == 0:
        raise ReadRejected("not_a_table")
    labels: list[str] = []
    for cell in header:
        if _is_blank_header(cell):
            raise ReadRejected("not_a_table")
        labels.append(_header_label(cell))
    return _dedupe_headers(labels)


def _apply_columns(frame: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    if len(frame) == 0:
        raise ReadRejected("no_data_rows")
    table = frame.copy()
    table.columns = columns
    table = table.replace("", pd.NA)
    table.index = pd.RangeIndex(stop=len(table))
    return table


def _is_blank_header(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _header_label(value: object) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _dedupe_headers(labels: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    columns: list[str] = []
    for label in labels:
        count = seen.get(label, 0) + 1
        seen[label] = count
        if count == 1:
            columns.append(label)
        else:
            columns.append(f"{label}_{count}")
    return columns
