"""Excel workbook for shared report blocks."""

from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook

from app.domain.report_model import ReportBlock


def report_xlsx(blocks: Sequence[ReportBlock]) -> bytes:
    """Write one sheet per block and return the workbook bytes.

    Row 1 is the block title. A block with a message writes that message on
    the next row and skips tables. Other blocks write each table as a header
    row followed by its rows, with a blank row between tables.
    """

    workbook = Workbook()
    for index, block in enumerate(blocks):
        if index == 0:
            sheet = workbook.active
            sheet.title = block.sheet
        else:
            sheet = workbook.create_sheet(title=block.sheet)
        sheet.cell(row=1, column=1, value=block.title)
        if block.message is not None:
            sheet.cell(row=2, column=1, value=block.message)
            continue
        row_index = 2
        for table_index, table in enumerate(block.tables):
            if table_index > 0:
                row_index += 1
            for column, header in enumerate(table.headers, start=1):
                sheet.cell(row=row_index, column=column, value=header)
            row_index += 1
            for data_row in table.rows:
                for column, cell in enumerate(data_row, start=1):
                    sheet.cell(row=row_index, column=column, value=cell)
                row_index += 1
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
